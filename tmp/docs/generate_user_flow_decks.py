from __future__ import annotations

import argparse
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from user_flow_pack_data import OUTPUT_ROOT, VERIFIED_ON, select_workflows
from user_flow_pptx import (
    COLORS,
    FONT_BODY,
    FONT_DISPLAY,
    accent_for,
    add_bullets,
    add_card,
    add_label_chip,
    add_picture_card,
    add_textbox,
    add_title_block,
    apply_background,
    make_presentation,
)


def add_info_card(slide, *, left: float, top: float, width: float, height: float, accent, label: str, body: str):
    add_card(slide, left, top, width, height, fill_color=COLORS["card"], line_color=COLORS["line"])
    add_label_chip(slide, label, left + 0.18, top + 0.16, accent=accent)
    add_textbox(
        slide,
        left + 0.16,
        top + 0.55,
        width - 0.32,
        height - 0.72,
        text=body,
        font_name=FONT_BODY,
        font_size=13.5,
        color=COLORS["ink"],
        auto_fit=True,
    )


def add_fact_stack(slide, *, left: float, top: float, width: float, height: float, accent, items: list[tuple[str, str]]):
    add_card(slide, left, top, width, height, fill_color=COLORS["card"], line_color=COLORS["line"])
    cursor = top + 0.18
    gap = 0.06
    available = height - 0.24
    row_height = available / max(1, len(items))

    for label, body in items:
        add_label_chip(slide, label, left + 0.16, cursor, accent=accent)
        add_textbox(
            slide,
            left + 0.16,
            cursor + 0.3,
            width - 0.32,
            row_height - 0.24,
            text=body,
            font_name=FONT_BODY,
            font_size=11.5,
            color=COLORS["ink"],
            auto_fit=True,
        )
        cursor += row_height + gap


def add_fact_grid(slide, *, left: float, top: float, width: float, height: float, accent, items: list[tuple[str, str]]):
    two_col = len(items) > 2
    if not two_col:
        for index, item in enumerate(items):
            add_info_card(
                slide,
                left=left + (index * ((width / len(items)) + 0.05)),
                top=top,
                width=(width / len(items)) - 0.05,
                height=height,
                accent=accent,
                label=item[0],
                body=item[1],
            )
        return

    col_w = (width - 0.18) / 2
    row_h = (height - 0.18) / 2
    for index, item in enumerate(items):
        row = index // 2
        col = index % 2
        add_info_card(
            slide,
            left=left + (col * (col_w + 0.18)),
            top=top + (row * (row_h + 0.18)),
            width=col_w,
            height=row_h,
            accent=accent,
            label=item[0],
            body=item[1],
        )


def add_cover_slide(prs, workflow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"{workflow.title} • Verified {VERIFIED_ON}")

    add_label_chip(slide, workflow.section, 0.6, 0.45, accent=accent)
    add_textbox(
        slide,
        0.62,
        1.0,
        6.0,
        1.4,
        text=workflow.title,
        font_name=FONT_DISPLAY,
        font_size=26,
        color=COLORS["ink"],
        bold=True,
        auto_fit=True,
    )
    add_textbox(
        slide,
        0.64,
        2.15,
        5.75,
        0.8,
        text=workflow.document_purpose,
        font_name=FONT_BODY,
        font_size=13,
        color=COLORS["muted"],
        auto_fit=True,
    )

    add_info_card(
        slide,
        left=0.65,
        top=3.1,
        width=2.55,
        height=1.18,
        accent=accent,
        label="Primary user",
        body=workflow.primary_user,
    )
    add_info_card(
        slide,
        left=3.42,
        top=3.1,
        width=2.9,
        height=1.18,
        accent=accent,
        label="Entry point",
        body=workflow.entry_point,
    )

    add_picture_card(
        slide,
        workflow.cover_image,
        left=7.1,
        top=0.82,
        width=5.55,
        height=5.95,
        crop=workflow.cover_crop,
        target_ratio=0.94,
        anchor="top",
    )

    add_textbox(
        slide,
        0.66,
        4.58,
        5.75,
        1.55,
        text=workflow.closing_text or "Use this deck as part of the full user-flow training pack.",
        font_name=FONT_BODY,
        font_size=12.5,
        color=COLORS["ink"],
        auto_fit=True,
    )
    return slide


def add_overview_slide(prs, workflow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"Workflow summary • {workflow.short_title}")
    add_title_block(
        slide,
        section=workflow.section,
        workflow_no=workflow.number,
        title="Workflow At A Glance",
        subtitle="Purpose, user, entry point, and the operational outcome of this workflow.",
    )

    add_info_card(
        slide,
        left=0.65,
        top=2.0,
        width=3.85,
        height=1.35,
        accent=accent,
        label="Purpose",
        body=workflow.document_purpose,
    )
    add_info_card(
        slide,
        left=4.66,
        top=2.0,
        width=3.85,
        height=1.35,
        accent=accent,
        label="Primary user",
        body=workflow.primary_user,
    )
    add_info_card(
        slide,
        left=8.67,
        top=2.0,
        width=4.0,
        height=1.35,
        accent=accent,
        label="Entry point",
        body=workflow.entry_point,
    )

    add_card(slide, 0.65, 3.65, 12.02, 2.72, fill_color=COLORS["card"], line_color=COLORS["line"])
    add_label_chip(slide, "Workflow summary", 0.83, 3.82, accent=accent)
    add_bullets(
        slide,
        0.87,
        4.22,
        11.6,
        1.9,
        workflow.workflow_summary,
        font_size=13.5,
        color=COLORS["ink"],
        bullet_color=accent,
    )
    return slide


def add_platform_map_slide(prs, workflow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"Role map • {workflow.short_title}")
    add_title_block(
        slide,
        section=workflow.section,
        workflow_no=workflow.number,
        title="Role Map And Start Points",
        subtitle="How the product hands off between roles across campaign setup, clinic onboarding, sharing, and caregiver playback.",
    )

    stage_y = 2.1
    stages = [
        ("Campaign setup", "Publisher"),
        ("Recruitment", "Field rep"),
        ("Clinic sharing", "Doctor or clinic staff"),
        ("Caregiver playback", "Caregiver"),
    ]
    stage_x = [0.7, 3.5, 6.3, 9.1]
    for index, ((stage, role), left) in enumerate(zip(stages, stage_x)):
        add_info_card(
            slide,
            left=left,
            top=stage_y,
            width=2.2,
            height=1.25,
            accent=accent,
            label=role,
            body=stage,
        )
        if index < len(stages) - 1:
            add_label_chip(slide, "→", left + 2.28, stage_y + 0.44, accent=accent)

    add_card(slide, 0.7, 3.95, 5.7, 2.65, fill_color=COLORS["card"], line_color=COLORS["line"])
    add_label_chip(slide, "Role map", 0.88, 4.12, accent=accent)
    role_bullets = [f"{role}: {description}" for role, description in workflow.role_map]
    add_bullets(
        slide,
        0.88,
        4.5,
        5.35,
        1.9,
        role_bullets,
        font_size=11.5,
        color=COLORS["ink"],
        bullet_color=accent,
    )

    add_card(slide, 6.62, 3.95, 6.03, 2.65, fill_color=COLORS["card"], line_color=COLORS["line"])
    add_label_chip(slide, "Workflow start points", 6.8, 4.12, accent=accent)
    start_bullets = [f"{label}: {entry}" for label, entry in workflow.start_points]
    add_bullets(
        slide,
        6.82,
        4.5,
        5.6,
        1.9,
        start_bullets,
        font_size=11.2,
        color=COLORS["ink"],
        bullet_color=accent,
    )
    return slide


def add_decision_slide(prs, workflow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"Decision points • {workflow.short_title}")
    add_title_block(
        slide,
        section=workflow.section,
        workflow_no=workflow.number,
        title="Decision Points",
        subtitle="Key branches or usage choices that change what the user should do next.",
    )

    card_width = (12.02 - (0.26 * (len(workflow.decision_points) - 1))) / len(workflow.decision_points)
    for index, (title, when_to_use, outcome) in enumerate(workflow.decision_points):
        left = 0.65 + (index * (card_width + 0.26))
        add_card(slide, left, 2.05, card_width, 4.55, fill_color=COLORS["card"], line_color=COLORS["line"])
        add_label_chip(slide, title, left + 0.16, 2.22, accent=accent)
        add_textbox(
            slide,
            left + 0.15,
            2.72,
            card_width - 0.3,
            1.35,
            text=when_to_use,
            font_name=FONT_BODY,
            font_size=12,
            color=COLORS["ink"],
            auto_fit=True,
        )
        add_label_chip(slide, "Outcome", left + 0.16, 4.2, accent=accent)
        add_textbox(
            slide,
            left + 0.15,
            4.58,
            card_width - 0.3,
            1.55,
            text=outcome,
            font_name=FONT_BODY,
            font_size=12,
            color=COLORS["muted"],
            auto_fit=True,
        )
    return slide


def step_fact_items(step):
    items = [
        ("What the user does", step.user_does),
        ("What the user sees", step.user_sees),
        ("Why it matters", step.why_it_matters),
        ("Expected result", step.expected_result),
    ]
    if step.trainer_notes:
        items.append(("Trainer note", step.trainer_notes))
    return items


def add_step_slide(prs, workflow, step):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"Step {step.number} • {workflow.short_title}")
    add_title_block(
        slide,
        section=workflow.section,
        workflow_no=workflow.number,
        title=f"Step {step.number}: {step.title}",
        subtitle=step.screenshot_caption,
    )

    items = step_fact_items(step)

    if step.deck_layout == "wide":
        add_picture_card(
            slide,
            step.screenshot_path,
            left=0.65,
            top=2.0,
            width=12.02,
            height=2.9,
            crop=step.deck_crop,
            target_ratio=3.55,
            anchor="top",
        )
        left_items = [items[0], items[2]] if len(items) >= 3 else items[:1]
        right_items = [items[1], items[3]] if len(items) >= 4 else items[1:2]
        add_fact_stack(slide, left=0.65, top=5.12, width=5.86, height=1.38, accent=accent, items=left_items)
        add_fact_stack(slide, left=6.81, top=5.12, width=5.86, height=1.38, accent=accent, items=right_items)
        if len(items) > 4:
            add_card(slide, 0.65, 6.58, 12.02, 0.42, fill_color=COLORS["card"], line_color=COLORS["line"])
            add_label_chip(slide, items[4][0], 0.82, 6.62, accent=accent)
            add_textbox(
                slide,
                2.48,
                6.61,
                9.95,
                0.22,
                text=items[4][1],
                font_name=FONT_BODY,
                font_size=10.5,
                color=COLORS["ink"],
                valign=MSO_ANCHOR.MIDDLE,
                auto_fit=True,
            )
    else:
        add_picture_card(
            slide,
            step.screenshot_path,
            left=0.65,
            top=2.0,
            width=6.6,
            height=4.9,
            crop=step.deck_crop,
            target_ratio=1.35,
            anchor="top",
        )
        add_fact_stack(slide, left=7.48, top=2.0, width=5.2, height=4.9, accent=accent, items=items[:5])
    return slide


def add_issues_slide(prs, workflow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"Trainer tips • {workflow.short_title}")
    add_title_block(
        slide,
        section=workflow.section,
        workflow_no=workflow.number,
        title="Trainer Tips And Common Issues",
        subtitle="The highest-value coaching points to cover during rollout or onboarding sessions.",
    )

    col_w = 5.85
    row_h = 1.68
    for index, (title, guidance) in enumerate(workflow.common_issues):
        row = index // 2
        col = index % 2
        left = 0.65 + (col * 6.17)
        top = 2.02 + (row * 1.88)
        add_info_card(
            slide,
            left=left,
            top=top,
            width=col_w,
            height=row_h,
            accent=accent,
            label=title,
            body=guidance,
        )

    add_card(slide, 0.65, 6.05, 12.02, 0.78, fill_color=COLORS["sand"], line_color=COLORS["line"])
    add_label_chip(slide, "Success markers", 0.83, 6.2, accent=accent)
    add_bullets(
        slide,
        2.35,
        6.16,
        10.0,
        0.54,
        workflow.success_criteria,
        font_size=11.5,
        color=COLORS["ink"],
        bullet_color=accent,
    )
    return slide


def add_closing_slide(prs, workflow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(workflow.section)
    apply_background(slide, workflow.section, footer_text=f"Workflow complete • {workflow.short_title}")
    add_label_chip(slide, "Workflow complete", 0.65, 0.55, accent=accent)
    add_textbox(
        slide,
        0.66,
        1.1,
        5.7,
        0.9,
        text=workflow.title,
        font_name=FONT_DISPLAY,
        font_size=24,
        color=COLORS["ink"],
        bold=True,
        auto_fit=True,
    )
    add_textbox(
        slide,
        0.68,
        2.0,
        5.7,
        0.72,
        text=workflow.closing_text or "Return to the index deck and continue with the next role-specific workflow.",
        font_name=FONT_BODY,
        font_size=12.5,
        color=COLORS["muted"],
        auto_fit=True,
    )

    add_card(slide, 0.68, 3.0, 5.7, 2.95, fill_color=COLORS["card"], line_color=COLORS["line"])
    add_label_chip(slide, "Success criteria", 0.86, 3.18, accent=accent)
    add_bullets(
        slide,
        0.86,
        3.58,
        5.28,
        1.98,
        workflow.success_criteria,
        font_size=12,
        color=COLORS["ink"],
        bullet_color=accent,
    )

    add_card(slide, 6.65, 1.0, 6.02, 4.95, fill_color=COLORS["card"], line_color=COLORS["line"])
    add_label_chip(slide, "Related manuals", 6.84, 1.18, accent=accent)
    manual_items = [item for item in workflow.related_documents if item.endswith(".md")][:5]
    if manual_items:
        add_bullets(
            slide,
            6.85,
            1.58,
            5.55,
            1.55,
            manual_items,
            font_size=12,
            color=COLORS["ink"],
            bullet_color=accent,
        )
    add_label_chip(slide, "Status", 6.84, 3.48, accent=accent)
    add_textbox(
        slide,
        6.85,
        3.86,
        5.4,
        1.12,
        text=workflow.status,
        font_name=FONT_BODY,
        font_size=12,
        color=COLORS["ink"],
        auto_fit=True,
    )
    return slide


def build_workflow_deck(workflow) -> Path:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    prs = make_presentation()
    add_cover_slide(prs, workflow)
    add_overview_slide(prs, workflow)

    if workflow.role_map or workflow.start_points:
        add_platform_map_slide(prs, workflow)
    elif workflow.decision_points:
        add_decision_slide(prs, workflow)

    for step in workflow.steps:
        add_step_slide(prs, workflow, step)

    if workflow.common_issues:
        add_issues_slide(prs, workflow)

    add_closing_slide(prs, workflow)

    target = OUTPUT_ROOT / workflow.deck_filename
    prs.save(target)
    print(f"Wrote {target}")
    return target


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate workflow PPT decks for the user-flow pack.")
    parser.add_argument(
        "--workflows",
        help="Comma-separated workflow numbers or slugs to regenerate. Defaults to the full pack.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for workflow in select_workflows(args.workflows):
        build_workflow_deck(workflow)


if __name__ == "__main__":
    main()
