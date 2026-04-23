from __future__ import annotations

import hashlib
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from user_flow_pack_data import DOCS_ROOT, REPO_ROOT


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"

COLORS = {
    "bg": (247, 245, 240),
    "card": (255, 255, 255),
    "ink": (39, 51, 64),
    "muted": (96, 111, 126),
    "line": (220, 227, 232),
    "navy": (52, 74, 163),
    "teal": (40, 178, 177),
    "mint": (220, 244, 242),
    "sand": (245, 236, 223),
    "coral": (236, 113, 83),
    "sage": (138, 180, 138),
    "plum": (95, 83, 142),
}

SECTION_ACCENTS = {
    "Platform and Campaign Setup": COLORS["teal"],
    "Clinic Onboarding and Sharing": COLORS["navy"],
    "Caregiver Experience": COLORS["sage"],
    "Internal Operations": COLORS["coral"],
}

CROP_CACHE = REPO_ROOT / "tmp" / "docs" / "generated-crops"


def rgb(value: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*value)


def accent_for(section: str) -> tuple[int, int, int]:
    return SECTION_ACCENTS.get(section, COLORS["teal"])


def make_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def add_card(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    fill_color: tuple[int, int, int] = COLORS["card"],
    line_color: tuple[int, int, int] = COLORS["line"],
    radius: MSO_AUTO_SHAPE_TYPE = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(radius, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_color)
    shape.line.color.rgb = rgb(line_color)
    shape.line.width = Pt(1)
    return shape


def add_label_chip(slide, text: str, left: float, top: float, *, accent: tuple[int, int, int]):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(max(1.0, 0.18 + (len(text) * 0.08))),
        Inches(0.34),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(accent)
    shape.line.color.rgb = rgb(accent)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(8)
    frame.margin_right = Pt(8)
    frame.margin_top = Pt(1)
    frame.margin_bottom = Pt(1)
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = FONT_BODY
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = rgb(COLORS["card"])
    p.alignment = PP_ALIGN.CENTER
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    text: str = "",
    font_name: str = FONT_BODY,
    font_size: float = 18,
    color: tuple[int, int, int] = COLORS["ink"],
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    auto_fit: bool = True,
    margin: float = 0.06,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(margin * 72)
    frame.margin_right = Pt(margin * 72)
    frame.margin_top = Pt(margin * 72)
    frame.margin_bottom = Pt(margin * 72)
    frame.vertical_anchor = valign
    if auto_fit:
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    p.alignment = align
    return box


def add_bullets(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    bullets: list[str],
    *,
    font_size: float = 16,
    color: tuple[int, int, int] = COLORS["ink"],
    bullet_color: tuple[int, int, int] | None = None,
    font_name: str = FONT_BODY,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Pt(4)
    frame.margin_right = Pt(4)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    for index, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.level = 0
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = f"• {bullet}"
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.color.rgb = rgb(bullet_color or color)
    return box


def apply_background(slide, section: str, footer_text: str = "") -> None:
    accent = accent_for(section)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(COLORS["bg"])

    top_bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0),
        Inches(0),
        SLIDE_WIDTH,
        Inches(0.12),
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = rgb(accent)
    top_bar.line.fill.background()

    orb_left = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(-0.75),
        Inches(5.5),
        Inches(2.5),
        Inches(2.5),
    )
    orb_left.fill.solid()
    orb_left.fill.fore_color.rgb = rgb(COLORS["sand"])
    orb_left.fill.transparency = 0.2
    orb_left.line.fill.background()

    orb_right = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(10.9),
        Inches(-0.55),
        Inches(2.9),
        Inches(2.9),
    )
    orb_right.fill.solid()
    orb_right.fill.fore_color.rgb = rgb(COLORS["mint"])
    orb_right.fill.transparency = 0.08
    orb_right.line.fill.background()

    footer = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0),
        Inches(7.18),
        SLIDE_WIDTH,
        Inches(0.32),
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = rgb((243, 240, 234))
    footer.line.fill.background()

    add_textbox(
        slide,
        0.45,
        7.2,
        6.5,
        0.18,
        text=footer_text or section,
        font_name=FONT_BODY,
        font_size=9,
        color=COLORS["muted"],
        bold=False,
        valign=MSO_ANCHOR.MIDDLE,
        auto_fit=False,
        margin=0.0,
    )


def add_title_block(
    slide,
    *,
    section: str,
    workflow_no: int | None,
    title: str,
    subtitle: str = "",
):
    accent = accent_for(section)
    if workflow_no is not None:
        add_label_chip(slide, f"Workflow {workflow_no:02d}", 0.55, 0.32, accent=accent)
    add_label_chip(slide, section, 2.0 if workflow_no is not None else 0.55, 0.32, accent=accent)

    add_textbox(
        slide,
        0.55,
        0.72,
        7.0,
        0.85,
        text=title,
        font_name=FONT_DISPLAY,
        font_size=24,
        color=COLORS["ink"],
        bold=True,
        auto_fit=True,
    )
    if subtitle:
        add_textbox(
            slide,
            0.57,
            1.45,
            7.1,
            0.48,
            text=subtitle,
            font_name=FONT_BODY,
            font_size=11.5,
            color=COLORS["muted"],
            auto_fit=True,
        )


def resolve_repo_path(relative_or_absolute: str) -> Path:
    path = Path(relative_or_absolute)
    if path.is_absolute():
        return path
    if relative_or_absolute.startswith("assets/"):
        return DOCS_ROOT / relative_or_absolute
    return REPO_ROOT / relative_or_absolute


def prepare_image(
    image_path: str | Path,
    *,
    crop: tuple[float, float, float, float] | None = None,
    target_ratio: float | None = None,
    anchor: str = "top",
) -> Path:
    source = resolve_repo_path(str(image_path))
    if not source.exists():
        raise FileNotFoundError(source)

    CROP_CACHE.mkdir(parents=True, exist_ok=True)
    key = f"{source}:{crop}:{target_ratio}:{anchor}"
    digest = hashlib.md5(key.encode("utf-8")).hexdigest()
    ext = source.suffix or ".png"
    target = CROP_CACHE / f"{source.stem}-{digest}{ext}"
    if target.exists():
        return target

    with Image.open(source) as image:
        width, height = image.size
        left, top, right, bottom = 0, 0, width, height

        if crop:
            c_left, c_top, c_width, c_height = crop
            left = int(width * c_left)
            top = int(height * c_top)
            right = int(width * min(1.0, c_left + c_width))
            bottom = int(height * min(1.0, c_top + c_height))
        elif target_ratio:
            current_ratio = width / height
            if abs(current_ratio - target_ratio) > 0.02:
                if current_ratio > target_ratio:
                    new_width = int(height * target_ratio)
                    left = max(0, int((width - new_width) / 2))
                    right = min(width, left + new_width)
                else:
                    new_height = int(width / target_ratio)
                    if anchor == "center":
                        top = max(0, int((height - new_height) / 2))
                    elif anchor == "bottom":
                        top = max(0, height - new_height)
                    else:
                        top = 0
                    bottom = min(height, top + new_height)

        cropped = image.crop((left, top, right, bottom))
        cropped.save(target)
    return target


def add_picture(
    slide,
    image_path: str | Path,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
):
    slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def add_picture_card(
    slide,
    image_path: str | Path,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    crop: tuple[float, float, float, float] | None = None,
    target_ratio: float | None = None,
    anchor: str = "top",
):
    frame = add_card(slide, left, top, width, height, fill_color=COLORS["card"], line_color=COLORS["line"])
    image_target = prepare_image(image_path, crop=crop, target_ratio=target_ratio, anchor=anchor)
    add_picture(
        slide,
        image_target,
        left=left + 0.08,
        top=top + 0.08,
        width=width - 0.16,
        height=height - 0.16,
    )
    return frame
